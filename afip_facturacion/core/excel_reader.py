"""Lectura del Excel con los comprobantes a emitir.

Una fila = un comprobante. La unica columna verdaderamente obligatoria es el
numero de documento del receptor: todo lo demas (fecha, periodo, importe,
descripcion) puede venir en el Excel o como valor por defecto para toda la
corrida.

Ese doble origen es a proposito. En el caso tipico -- la cuota social de un
mes -- la fecha, el periodo y el importe son identicos para los cuarenta
socios, y obligar a repetirlos cuarenta veces es una fuente de errores de
tipeo. Entonces se pasan una vez por linea de comandos y el Excel queda con
lo unico que cambia: quien. Pero si un socio paga una cuota distinta, alcanza
con poner esa columna en el Excel y esa fila la usa.

Los nombres de columna se buscan sin distinguir mayusculas, tildes ni
espacios de mas, porque un Excel escrito por una persona dice "DNI", "dni" o
"Nro. Documento" segun el dia.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from .models import ComprobanteRequest, resolver_condicion_iva

#: Nombre canonico -> los alias que puede tener esa columna en el Excel.
#: Se comparan ya normalizados (sin tildes, minusculas, sin espacios extra).
ALIAS_COLUMNAS = {
    "documento": ["documento", "dni", "nro documento", "numero de documento",
                  "nro doc", "n documento", "n doc", "nro", "numero documento",
                  "documento nro", "cuit", "doc"],
    "nombre": ["nombre", "socio", "apellido y nombre", "nombre y apellido",
               "razon social", "cliente"],
    "importe": ["importe", "monto", "precio", "total", "valor", "cuota"],
    "fecha": ["fecha", "fecha comprobante", "fecha del comprobante"],
    "periodo_desde": ["periodo desde", "desde", "servicio desde"],
    "periodo_hasta": ["periodo hasta", "hasta", "servicio hasta"],
    "vencimiento_pago": ["vencimiento", "vto pago", "vto para el pago",
                         "vencimiento de pago", "vto"],
    "descripcion": ["descripcion", "concepto", "detalle",
                    "descripcion del servicio"],
    "condicion_iva": ["condicion iva", "condicion frente al iva", "iva"],
    "tipo_documento": ["tipo documento", "tipo de documento", "tipo doc"],
}

TIPOS_DOCUMENTO = {"dni": 96, "cuit": 80, "consumidor final": 99, "cf": 99}


class ExcelError(Exception):
    """El Excel no se pudo leer, o le falta algo imprescindible."""


@dataclass
class FilaInvalida:
    """Una fila que no se pudo interpretar, con el motivo.

    Se junta en vez de cortar la corrida: si la fila 12 tiene el importe mal
    escrito, conviene emitir las otras 39 y reportar esa, no frenar todo.
    """

    numero: int
    motivo: str


@dataclass
class Defaults:
    """Valores que aplican a toda la corrida si la fila no los trae."""

    fecha: date | None = None
    periodo_desde: date | None = None
    periodo_hasta: date | None = None
    vencimiento_pago: date | None = None
    importe: float | None = None
    descripcion: str = ""


#: Caracteres que se tratan como separadores al comparar encabezados. Los
#: indicadores ordinales entran aca porque en castellano "numero" se abrevia
#: "N°" o "Nº" indistintamente, y ninguno de los dos es un acento que la
#: normalizacion Unicode saque sola: sin esto, "N° Documento" no matchearia
#: con el alias "n documento".
_SEPARADORES = ".", "_", "°", "º", "#", "-", "/"


def _normalizar(texto: Any) -> str:
    """Minusculas, sin tildes y sin espacios de mas, para comparar encabezados.

    El orden de los dos pasos importa: los separadores se reemplazan ANTES
    de normalizar los acentos. Si se hiciera al reves, NFKD convierte "º" en
    la letra "o" y "Nº Doc" terminaria como "no doc" en vez de "n doc".
    """
    if texto is None:
        return ""
    limpio = str(texto)
    for separador in _SEPARADORES:
        limpio = limpio.replace(separador, " ")
    sin_tildes = unicodedata.normalize("NFKD", limpio)
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return " ".join(sin_tildes.lower().split())


def _mapear_columnas(encabezados: list[Any]) -> dict[str, int]:
    """Encabezados del Excel -> indice de columna, por nombre canonico."""
    normalizados = [_normalizar(e) for e in encabezados]
    mapa: dict[str, int] = {}
    for canonico, alias in ALIAS_COLUMNAS.items():
        for indice, encabezado in enumerate(normalizados):
            if encabezado and encabezado in alias:
                mapa[canonico] = indice
                break
    return mapa


def _a_fecha(valor: Any, campo: str) -> date | None:
    """Convierte una celda a fecha, aceptando lo que escriba una persona."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(f"{campo}: no se entiende la fecha {valor!r} (usar dd/mm/aaaa)")


def _a_importe(valor: Any) -> float | None:
    """Convierte una celda a importe, tolerando '$', puntos de miles y comas.

    El caso ambiguo es "21.500": puede leerse como veintiun mil quinientos
    (punto de miles, la convencion argentina) o como 21 con 50. Se resuelve
    mirando cuantos digitos siguen al ultimo punto: si son exactamente tres,
    es separador de miles; si son uno o dos, es decimal. La regla funciona
    porque un importe en pesos no lleva tres decimales, asi que "21.500"
    con tres digitos atras solo puede ser miles.

    Nada de esto aplica cuando la celda es numerica de verdad: ahi openpyxl
    ya devuelve un float y no hay nada que interpretar. Este camino es solo
    para celdas escritas como texto.
    """
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if "," in texto:
        # Con coma presente no hay ambiguedad: coma decimal, punto de miles.
        texto = texto.replace(".", "").replace(",", ".")
    elif "." in texto:
        ultimo_grupo = texto.rsplit(".", 1)[1]
        if len(ultimo_grupo) == 3:
            texto = texto.replace(".", "")
    try:
        return float(texto)
    except ValueError:
        raise ValueError(f"importe: no se entiende el valor {valor!r}") from None


def _a_tipo_documento(valor: Any) -> int:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return TIPOS_DOCUMENTO["dni"]
    if isinstance(valor, (int, float)):
        return int(valor)
    texto = _normalizar(valor)
    if texto.isdigit():
        return int(texto)
    if texto not in TIPOS_DOCUMENTO:
        raise ValueError(
            f"tipo de documento desconocido: {valor!r} "
            f"(usar {', '.join(sorted(TIPOS_DOCUMENTO))})"
        )
    return TIPOS_DOCUMENTO[texto]


def leer_comprobantes(
    ruta: str,
    hoja: str | None = None,
    defaults: Defaults | None = None,
) -> tuple[list[ComprobanteRequest], list[FilaInvalida]]:
    """Lee el Excel y devuelve (comprobantes validos, filas con problemas)."""
    defaults = defaults or Defaults()
    try:
        libro = load_workbook(ruta, data_only=True, read_only=True)
    except FileNotFoundError:
        raise ExcelError(f"no se encontro el archivo {ruta}") from None
    except Exception as exc:
        raise ExcelError(f"no se pudo abrir {ruta}: {exc}") from exc

    try:
        if hoja:
            if hoja not in libro.sheetnames:
                raise ExcelError(
                    f"el Excel no tiene una hoja llamada {hoja!r}. "
                    f"Tiene: {', '.join(libro.sheetnames)}"
                )
            pagina = libro[hoja]
        else:
            pagina = libro.active

        filas = pagina.iter_rows(values_only=True)
        try:
            encabezados = list(next(filas))
        except StopIteration:
            raise ExcelError("la hoja esta vacia") from None

        columnas = _mapear_columnas(encabezados)
        if "documento" not in columnas:
            raise ExcelError(
                "no se encontro la columna del documento. Se busco alguna de: "
                + ", ".join(ALIAS_COLUMNAS["documento"])
                + f". Encabezados encontrados: {[e for e in encabezados if e]}"
            )

        comprobantes: list[ComprobanteRequest] = []
        invalidas: list[FilaInvalida] = []

        for numero_fila, fila in enumerate(filas, start=2):
            # Una fila en blanco se saltea, no corta la lectura: los Excel
            # reales tienen renglones vacios de separacion en el medio, y
            # cortar ahi dejaria socios sin facturar sin avisar nada.
            if all(celda is None or str(celda).strip() == "" for celda in fila):
                continue
            try:
                comprobantes.append(_fila_a_comprobante(fila, columnas, defaults, numero_fila))
            except (ValueError, KeyError) as exc:
                invalidas.append(FilaInvalida(numero=numero_fila, motivo=str(exc)))

        return comprobantes, invalidas
    finally:
        libro.close()


def _celda(fila: tuple, columnas: dict[str, int], nombre: str) -> Any:
    indice = columnas.get(nombre)
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def _fila_a_comprobante(
    fila: tuple,
    columnas: dict[str, int],
    defaults: Defaults,
    numero_fila: int,
) -> ComprobanteRequest:
    documento = _celda(fila, columnas, "documento")
    if documento is None or not str(documento).strip():
        raise ValueError("falta el numero de documento")
    documento = "".join(ch for ch in str(documento) if ch.isdigit())
    if not documento:
        raise ValueError("el documento no tiene digitos")

    importe = _a_importe(_celda(fila, columnas, "importe"))
    if importe is None:
        importe = defaults.importe
    if importe is None:
        raise ValueError("falta el importe (ni en el Excel ni en --importe)")

    fecha = _a_fecha(_celda(fila, columnas, "fecha"), "fecha") or defaults.fecha
    if fecha is None:
        raise ValueError("falta la fecha (ni en el Excel ni en --fecha)")

    descripcion = _celda(fila, columnas, "descripcion")
    descripcion = str(descripcion).strip() if descripcion else defaults.descripcion

    return ComprobanteRequest(
        doc_nro=documento,
        importe=importe,
        fecha=fecha,
        periodo_desde=(_a_fecha(_celda(fila, columnas, "periodo_desde"), "periodo desde")
                       or defaults.periodo_desde),
        periodo_hasta=(_a_fecha(_celda(fila, columnas, "periodo_hasta"), "periodo hasta")
                       or defaults.periodo_hasta),
        vencimiento_pago=(_a_fecha(_celda(fila, columnas, "vencimiento_pago"), "vencimiento")
                          or defaults.vencimiento_pago),
        descripcion=descripcion,
        doc_tipo=_a_tipo_documento(_celda(fila, columnas, "tipo_documento")),
        condicion_iva=resolver_condicion_iva(_celda(fila, columnas, "condicion_iva")),
        nombre=str(_celda(fila, columnas, "nombre") or "").strip(),
        fila_excel=numero_fila,
    )
