"""La linea de comandos: argumentos, config, dry-run y salida."""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook

from afip_facturacion.cli import _fecha_arg, cargar_emisor, construir_parser, main


@pytest.fixture
def excel(tmp_path):
    libro = Workbook()
    pagina = libro.active
    pagina.title = "Julio"
    pagina.append(["Socio", "DNI", "Importe"])
    pagina.append(["Ana Perez", "30111222", 21500])
    pagina.append(["Luis Gomez", "30111223", None])
    ruta = tmp_path / "cuotas.xlsx"
    libro.save(ruta)
    return str(ruta)


@pytest.fixture
def config(tmp_path):
    ruta = tmp_path / "config.py"
    ruta.write_text(
        'CUIT = "20-11122233-4"\n'
        "PUNTO_VENTA = 3\n"
        "CBTE_TIPO = 15\n"
        "CONCEPTO = 2\n"
    )
    return str(ruta)


# ---------------------------------------------------------------------------
# Argumentos
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("texto,esperado", [
    ("27/07/2026", date(2026, 7, 27)),
    ("2026-07-27", date(2026, 7, 27)),
    ("27-07-2026", date(2026, 7, 27)),
])
def test_parsea_fechas_de_la_linea_de_comandos(texto, esperado):
    assert _fecha_arg(texto) == esperado


def test_rechaza_una_fecha_invalida():
    import argparse
    with pytest.raises(argparse.ArgumentTypeError):
        _fecha_arg("31 de julio")


def test_produccion_es_opt_in():
    """Por defecto va a homologacion: emitir de verdad tiene que ser explicito."""
    args = construir_parser().parse_args(["--excel", "x.xlsx"])
    assert args.produccion is False


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def test_carga_la_config_del_emisor(config):
    emisor = cargar_emisor(config, produccion=False)
    assert emisor.cuit == "20111222334"       # normalizado, sin guiones
    assert emisor.punto_venta == 3
    assert emisor.cbte_tipo == 15
    assert emisor.produccion is False


def test_el_flag_produccion_pisa_la_config(config):
    assert cargar_emisor(config, produccion=True).produccion is True


def test_config_inexistente_explica_que_hacer(tmp_path):
    with pytest.raises(SystemExit) as error:
        cargar_emisor(str(tmp_path / "no_existe.py"), produccion=False)
    assert "config.example.py" in str(error.value)


def test_config_incompleta_dice_que_falta(tmp_path):
    ruta = tmp_path / "config.py"
    ruta.write_text('CUIT = "20111222334"\n')      # sin PUNTO_VENTA
    with pytest.raises(SystemExit) as error:
        cargar_emisor(str(ruta), produccion=False)
    assert "PUNTO_VENTA" in str(error.value)


def test_config_con_cuit_invalido(tmp_path):
    ruta = tmp_path / "config.py"
    ruta.write_text('CUIT = "123"\nPUNTO_VENTA = 1\n')
    with pytest.raises(SystemExit) as error:
        cargar_emisor(str(ruta), produccion=False)
    assert "CUIT invalido" in str(error.value)


# ---------------------------------------------------------------------------
# Corrida completa (sin tocar AFIP)
# ---------------------------------------------------------------------------

def test_dry_run_no_llama_a_afip(excel, config, capsys):
    codigo = main([
        "--excel", excel, "--hoja", "Julio", "--config", config,
        "--fecha", "27/07/2026", "--importe", "21500", "--dry-run",
    ])
    salida = capsys.readouterr().out
    assert codigo == 0
    assert "no se llamo a AFIP" in salida
    assert "Ana Perez" in salida
    assert "homologacion" in salida


def test_dry_run_aplica_el_importe_por_defecto(excel, config, capsys):
    """La segunda fila no trae importe: tiene que tomar el de --importe."""
    main(["--excel", excel, "--hoja", "Julio", "--config", config,
          "--fecha", "27/07/2026", "--importe", "21500", "--dry-run"])
    salida = capsys.readouterr().out
    assert salida.count("21,500.00") == 2


def test_reporta_las_filas_invalidas(tmp_path, config, capsys):
    libro = Workbook()
    libro.active.append(["Socio", "DNI"])
    libro.active.append(["Sin documento", None])
    libro.active.append(["Ana", "30111222"])
    ruta = tmp_path / "roto.xlsx"
    libro.save(ruta)

    main(["--excel", str(ruta), "--config", config,
          "--fecha", "27/07/2026", "--importe", "1000", "--dry-run"])
    salida = capsys.readouterr().out
    assert "1 fila(s) con problemas" in salida
    assert "fila 2" in salida


def test_excel_inexistente_devuelve_codigo_de_error(config, capsys):
    codigo = main(["--excel", "/no/existe.xlsx", "--config", config, "--dry-run"])
    assert codigo == 2
    assert "Error leyendo el Excel" in capsys.readouterr().err


def test_excel_sin_filas_utiles(tmp_path, config, capsys):
    libro = Workbook()
    libro.active.append(["Socio", "DNI"])
    ruta = tmp_path / "vacio.xlsx"
    libro.save(ruta)

    codigo = main(["--excel", str(ruta), "--config", config,
                   "--fecha", "27/07/2026", "--importe", "1000"])
    assert codigo == 0
    assert "No hay nada para emitir" in capsys.readouterr().out
